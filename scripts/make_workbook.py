# -*- coding: utf-8 -*-
"""把导出的 CSV 汇总成一个多工作表的 Excel 工作簿。"""
import os, csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TBL = os.path.join(ROOT, "docs", "tables")
OUT = os.path.join(ROOT, "诗海图谱_数据表.xlsx")

SHEETS = [
    ("01_分期统计.csv", "分期统计"),
    ("02_体裁统计.csv", "体裁统计"),
    ("03_诗人产量Top500.csv", "产量Top500"),
    ("04_高频词Top300.csv", "高频词"),
    ("05_分期特征词.csv", "分期特征词"),
    ("06_主题模型_词分布.csv", "主题词分布"),
    ("07_地名统计.csv", "地名统计"),
    ("08_诗人档案Top1200.csv", "诗人档案"),
]

wb = openpyxl.Workbook()
wb.remove(wb.active)
head_fill = PatternFill("solid", fgColor="B23A2F")
head_font = Font(color="FFFFFF", bold=True, name="Microsoft YaHei", size=10)
body_font = Font(name="Times New Roman", size=10)
thin = Side(style="thin", color="DDD2BF")

for fname, title in SHEETS:
    path = os.path.join(TBL, fname)
    if not os.path.exists(path):
        print("skip", fname)
        continue
    ws = wb.create_sheet(title)
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    for r in rows[:5000]:
        ws.append(r)
    for c in range(1, len(rows[0]) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = body_font
            cell.border = Border(bottom=thin)
    widths = {}
    for row in ws.iter_rows(min_row=1, max_row=min(400, ws.max_row)):
        for cell in row:
            v = str(cell.value or "")
            w = sum(2 if ord(ch) > 127 else 1 for ch in v)
            widths[cell.column] = max(widths.get(cell.column, 8), min(w + 3, 46))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

wb.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")
